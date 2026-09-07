from contextlib import closing
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
import os
import re
import tempfile

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from .archive import Archive, tables
from .common import ArchiveError, checkpoint, connect_ro
from .messages import display_text, quoted_images


def clean(text):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]', '', str(text))


@dataclass
class ExportResult:
    path: Path
    messages: int
    pictures: int
    missing: int


def add_embedded_picture(doc, data, description, status):
    image = BytesIO(data)
    with Image.open(image) as picture:
        width = min(15, 18 * picture.width / picture.height)
    image.seek(0)
    paragraph = doc.add_paragraph()
    picture = paragraph.add_run().add_picture(image, width=Cm(width))
    picture._inline.docPr.set('descr', clean(description))
    paragraph.paragraph_format.keep_together = True
    if status != '原图':
        doc.add_paragraph('[' + clean(status) + ']').runs[0].font.size = Pt(8)


def reference_picture(con, conversation, reference, has_pictures):
    sid = reference['server_id']
    if not has_pictures or not sid.isdigit() or not int(sid):
        return None, '存档没有可匹配的原图片消息'
    # A reference must match the same conversation. Never use nearby timestamps
    # or the first similar image as a substitute for an exact message identity.
    rows = con.execute("""SELECT id FROM messages
        WHERE conversation=? AND server_id=? AND kind='图片'""", (conversation, sid)).fetchall()
    if len(rows) != 1:
        return None, '原图片消息不在存档中' if not rows else '原图片消息存在多个候选，无法确定对应图片'
    row = con.execute('SELECT data,status FROM pictures WHERE message_id=?', (rows[0][0],)).fetchone()
    return (row[0], row[1]) if row else (None, '原图片尚未还原')


def export_docx(archive: Archive, chat_id: str, target: Path, progress=lambda s: None, cancel=None,
                overwrite=False) -> ExportResult:
    chat = next((c for c in archive.conversations() if c['id'] == chat_id), None)
    if not chat:
        raise ArchiveError('请选择一个联系人或群聊。')
    if not chat['count']:
        raise ArchiveError('这个会话没有消息。')
    target = target.resolve()
    if target.suffix.lower() != '.docx':
        raise ArchiveError('导出文件扩展名必须为 .docx。')
    if target.exists() and not overwrite:
        raise ArchiveError('文件已存在，请使用不同的文件名。')
    source = Path(archive.meta['media_root']).resolve()
    if target.is_relative_to(source) or target == archive.path:
        raise ArchiveError('请将 Word 文件保存到微信源账号目录之外。')
    target.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.8)
    section.left_margin = section.right_margin = Cm(2.1)
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(10)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.2
    for style_name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
    doc.add_heading(clean(chat['name']) + ' 聊天记录', 0)
    doc.add_paragraph('聊天对象：' + clean(chat['name']) + '（' + clean(chat['id']) + '）')
    doc.add_paragraph(f'共 {chat["count"]:,} 条消息 · 时间为本机本地时区 · 导出于 {datetime.now():%Y-%m-%d %H:%M}')
    if archive.meta.get('demo'):
        doc.add_paragraph('演示文档：以下消息和图片均为合成测试数据。')
    summary = doc.add_paragraph()
    footer = section.footer.paragraphs[0]
    footer.add_run('拾语 · 本地微信记录    |    ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    messages = pictures = missing = 0
    previous_day = ''
    with closing(connect_ro(archive.path)) as con:
        has_pictures = 'pictures' in tables(con)
        for msg in archive.messages(chats=[chat_id], cancel=cancel):
            checkpoint(cancel)
            stamp = datetime.fromtimestamp(msg['time'])
            day = stamp.strftime('%Y-%m-%d')
            if day != previous_day:
                doc.add_heading(day, 2)
                previous_day = day
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            if chat['is_group']:
                destination = chat['name']
            else:
                destination = chat['name'] if msg['is_self'] else '我' if archive.meta.get('self_id') else chat['name']
            sender = msg['sender']
            if msg['sender_id'] and not msg['is_self']:
                sender += ' (' + msg['sender_id'] + ')'
            run = p.add_run(clean(f'{stamp:%H:%M:%S}   {sender} → {destination}'))
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string('315B46')
            if msg['kind'] == '图片':
                row = con.execute('SELECT data,status FROM pictures WHERE message_id=?', (msg['id'],)).fetchone() if has_pictures else None
                if row and row[0]:
                    add_embedded_picture(doc, row[0], f'{stamp:%Y-%m-%d %H:%M:%S} {sender}发送的图片', row[1])
                    pictures += 1
                else:
                    reason = row[1] if row else '存档未包含此图片'
                    doc.add_paragraph('[图片未导出：' + clean(reason) + ']')
                    missing += 1
            else:
                doc.add_paragraph(clean(display_text(msg)))
                if msg['kind'] in ('引用', '分享'):
                    for reference in quoted_images(msg.get('raw', '')):
                        checkpoint(cancel)
                        data, status = reference_picture(con, chat_id, reference, has_pictures)
                        if data:
                            caption = doc.add_paragraph('引用图片 · ' + clean(reference['sender']))
                            caption.paragraph_format.keep_with_next = True
                            caption.runs[0].font.size = Pt(8)
                            add_embedded_picture(doc, data, f'引用 {reference["sender"]} 的图片', status)
                            pictures += 1
                        else:
                            doc.add_paragraph('[引用图片未导出：' + clean(status) + ']')
                            missing += 1
            messages += 1
            if messages % 100 == 0:
                progress(f'正在生成 Word · {messages:,}/{chat["count"]:,} 条 · 已嵌入 {pictures} 张图片')
    summary.text = f'已嵌入 {pictures} 张图片' + (f'；{missing} 张图片无法还原，原因标注在对应消息处。' if missing else '。')
    checkpoint(cancel)
    fd, name = tempfile.mkstemp(prefix='.docx-', suffix='.docx', dir=target.parent)
    os.close(fd)
    temporary = Path(name)
    try:
        doc.save(temporary)
        checkpoint(cancel)
        if overwrite:
            os.replace(temporary, target)
        else:
            temporary.rename(target)
    finally:
        temporary.unlink(missing_ok=True)
    progress(f'导出完成 · {messages:,} 条消息 · {pictures} 张图片')
    return ExportResult(target, messages, pictures, missing)
